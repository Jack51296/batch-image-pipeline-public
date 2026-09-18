# -*- coding: utf-8 -*-
"""从网盘的《光照种类-提示词总表.xlsx》筛出进度含"通过"的可用组合，导出 CSV。

用法: python filter_pass.py   （总表更新后重跑一次即可刷新包根的 提示词总表-通过清单.csv）
依赖: pip install openpyxl
"""
import sys, os, csv
sys.stdout.reconfigure(encoding="utf-8")
from openpyxl import load_workbook

XLSX = r"\\10.0.0.11\share\frank\光影图生图\光照种类-提示词总表.xlsx"
PKG = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(PKG, "提示词总表-通过清单.csv")

wb = load_workbook(XLSX, read_only=True, data_only=True)
ws = wb["提示词总表"]
rows = list(ws.iter_rows(values_only=True))

passed = []
for r in rows[1:]:
    if not r or not r[0]:
        continue
    src, status, tgt = str(r[0]).strip(), str(r[2] or "").strip(), str(r[3] or "").strip()
    has_prompt = bool(str(r[4] or "").strip())
    if "通过" in status:
        passed.append((src, tgt, status.replace("\n", " ")[:30],
                       "有prompt" if has_prompt else "无prompt"))

with open(OUT, "w", newline="", encoding="utf-8-sig") as f:
    w = csv.writer(f)
    w.writerow(["源光效", "生成目标", "进度", "prompt状态"])
    w.writerows(passed)

print(f"总行数 {len(rows)-1}，通过 {len(passed)} 行 -> {OUT}")
for x in passed:
    print(" | ".join(x))
