#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""读取交付 CSV + align_pair_details.csv，生成 Labkit 五列表（六类反向命名）。"""
from __future__ import print_function

import argparse
import csv
import os


READABLE_PREFIX = "mmu:aiplatform-temp:label_data/20260827_AI六妆转秋冬妆_labkit/"
REVERSE_LABELS = {
    "混血妆转去妆": "AI去妆转混血妆",
    "混血妆转Y2K": "AIY2K妆转混血妆",
    "混血妆转清透氧气": "AI清透氧气妆转混血妆",
    "混血妆转泰系": "AI泰系妆转混血妆",
    "混血妆转氛围感": "AI氛围感妆转混血妆",
    "混血妆转千金妆": "AI千金妆转混血妆",
}


def arguments():
    p = argparse.ArgumentParser(description="生成 Labkit 导入表")
    p.add_argument("--delivery-csv", required=True, help="原交付 CSV，需含 id、标签")
    p.add_argument("--align-csv", required=True, help="align_pair_details.csv")
    p.add_argument("--subdir", required=True, help="BlobStore 可读根前缀下的子目录")
    p.add_argument("--output", required=True, help="输出 .xlsx 或 .csv")
    return p.parse_args()


def read_csv(path):
    with open(path, "r", encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def build_rows(delivery_csv, align_csv, subdir):
    labels = {r["id"].strip(): r["标签"].strip() for r in read_csv(delivery_csv)}
    prefix = READABLE_PREFIX + subdir.strip("/") + "/"
    rows = []
    for index, pair in enumerate(read_csv(align_csv), 1):
        pair_id = pair["id"].strip()
        old_label = labels.get(pair_id)
        category = REVERSE_LABELS.get(old_label)
        if not category:
            raise RuntimeError("找不到或未配置类别: %s / %s" % (pair_id, old_label))
        image1 = os.path.basename(pair["input_image_path"].strip())
        image2 = os.path.basename(pair["output_image_path"].strip())
        rows.append([index, category, prefix + image1, prefix + image2, "BlobStore"])
    return rows


def write_csv(path, rows):
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["dataIndex", "category", "image1", "image2", "source"])
        w.writerows(rows)


def write_xlsx(path, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except ImportError:
        raise SystemExit("缺少 openpyxl；可把 --output 改成 .csv，或安装 openpyxl")
    wb = Workbook()
    ws = wb.active
    ws.title = "工作表1"
    ws.append(["dataIndex", "category", "image1", "image2", "source"])
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="E5E7EB")
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 28
    ws.column_dimensions["C"].width = 90
    ws.column_dimensions["D"].width = 90
    ws.column_dimensions["E"].width = 16
    info = wb.create_sheet("info")
    info["A200"] = "快手内部文档请勿外传"
    wb.save(path)


def main():
    args = arguments()
    rows = build_rows(args.delivery_csv, args.align_csv, args.subdir)
    if args.output.lower().endswith(".xlsx"):
        write_xlsx(args.output, rows)
    elif args.output.lower().endswith(".csv"):
        write_csv(args.output, rows)
    else:
        raise SystemExit("--output 必须以 .xlsx 或 .csv 结尾")
    counts = {}
    for row in rows:
        counts[row[1]] = counts.get(row[1], 0) + 1
    print("OUTPUT", args.output)
    print("ROWS", len(rows))
    for name in sorted(counts):
        print(name, counts[name])
    print("FIRST", rows[0])
    print("LAST", rows[-1])


if __name__ == "__main__":
    main()
