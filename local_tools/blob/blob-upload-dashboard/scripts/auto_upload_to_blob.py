# -*- coding: utf-8 -*-
"""
自动上传 upload_content 下的交付文件到 BS3（bucket-b / Blender_data/）。

目录映射规则（本地相对路径 = 线上 key 相对路径）：
    D:\\Blender白膜渲染\\视频可视化展示\\视频可视化展示\\upload_content\\0729\\target_video\\target_video_01.mp4
        -> bucket = bucket-b
           key    = Blender_data/0729/target_video/target_video_01.mp4
           url    = https://blobstore.example.com/bucket-b/Blender_data/0729/target_video/target_video_01.mp4?x-bs-client-force=true&ts=<毫秒>

常用命令：
    python auto_upload_to_blob.py --once                # 扫描一次，上传新增/改动文件
    python auto_upload_to_blob.py --watch               # 常驻监听，丢文件进去就自动传
    python auto_upload_to_blob.py --once --dir 0805     # 只处理 upload_content\\0805
    python auto_upload_to_blob.py --once --dry-run      # 只看要传什么，不真传
    python auto_upload_to_blob.py --once --fill-xlsx    # 传完顺手把 URL 回填进当天的 xlsx
    python auto_upload_to_blob.py --once --repair       # 核对线上大小，只重传损坏/缺失的文件
"""
import argparse
import csv
import fnmatch
import json
import mimetypes
import os
import sys
import threading
import time
from urllib.parse import quote

import boto3
import botocore
from botocore.config import Config
from boto3.s3.transfer import TransferConfig

# ----------------------------------------------------------------------------
# 配置
#
# 每个人的本地路径和线上前缀都不一样，所以这些值放在 blob_config.json 里，
# 不要改这个脚本。找配置的顺序：环境变量 BLOB_CONFIG → 脚本旁边 → 当前目录。
# 找不到就用下面的默认值，等于老行为。
# ----------------------------------------------------------------------------
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CONFIG_NAME = "blob_config.json"

DEFAULTS = {
    "local_root": os.path.join(BASE_DIR, "upload_content"),
    "bucket": "bucket-b",
    "key_prefix": "Blender_data",
    # Windows 办公网只能走 corp 域名；blobstore.example.internal 只在公司内网机器（Linux 开发机）上可达
    "endpoint": "https://blobstore.example.com",
    "public_host": "https://blobstore.example.com",
    "service_header": "grpc_onlineEarningGenRpcService",
}


def find_config():
    for cand in (os.environ.get("BLOB_CONFIG"),
                 os.path.join(BASE_DIR, CONFIG_NAME),
                 os.path.join(os.getcwd(), CONFIG_NAME)):
        if cand and os.path.isfile(cand):
            return cand
    return None


def load_config(path=None):
    cfg = dict(DEFAULTS)
    path = path or find_config()
    if not path:
        return cfg, None
    try:
        with open(path, "r", encoding="utf-8") as f:
            user = json.load(f)
    except Exception as e:
        print("[!] 配置读不了 %s：%s，先用默认值" % (path, e))
        return cfg, None
    for k, v in user.items():
        if k.startswith("_") or v in ("", None):
            continue
        cfg[k] = v
    cfg["local_root"] = os.path.abspath(os.path.expanduser(cfg["local_root"]))
    cfg["key_prefix"] = str(cfg["key_prefix"]).strip("/")
    return cfg, path


CONFIG, CONFIG_PATH = load_config()

LOCAL_ROOT = CONFIG["local_root"]
BUCKET = CONFIG["bucket"]
KEY_PREFIX = CONFIG["key_prefix"]
ENDPOINT = CONFIG["endpoint"]
PUBLIC_HOST = CONFIG["public_host"]
SERVICE_HEADER = CONFIG["service_header"]

STATE_FILE = os.path.join(LOCAL_ROOT, ".blob_upload_state.json")
# URL 清单按批次命名：0806/0806_blob_urls.csv、0729/0729_blob_urls.csv …
# 早期版本每个批次都叫 _blob_urls.csv，在文件选择框里根本分不清是哪天的
URL_CSV_SUFFIX = "_blob_urls.csv"
LEGACY_URL_CSV = "_blob_urls.csv"
# 看板托管到 BS3 后没有后端接口可用，靠这份索引自描述线上有什么
INDEX_NAME = "dashboard_index.json"

MEDIA_EXT = (".mp4", ".mov", ".webm", ".png", ".jpg", ".jpeg", ".webp")
ALLOW_EXT = {".mp4", ".mov", ".webm", ".png", ".jpg", ".jpeg", ".webp",
             ".csv", ".xlsx", ".json", ".txt", ".md"}
SKIP_DIRS = {"frames", "_logs", "_build", "__pycache__", ".git", "_archive"}
# 清单 csv 和索引 json 由 publish 阶段显式上传，不能让扫描器碰：
# 它们的内容依赖上传结果，被扫进来会形成「上传→内容变→再上传」的死循环
SKIP_PATTERNS = ["~$*", ".*", "*.tmp", "*.part", "*.crdownload", "*.lock",
                 "*" + URL_CSV_SUFFIX,      # 也能盖住旧的 _blob_urls.csv
                 INDEX_NAME]

STABLE_SECONDS = 5      # 文件最后修改超过这么久才认为写完了（避免传到半截的 mp4）
WATCH_INTERVAL = 10     # 监听模式扫描间隔（秒）

TRANSFER_CONFIG = TransferConfig(
    multipart_threshold=64 * 1024 * 1024,
    multipart_chunksize=16 * 1024 * 1024,
    max_concurrency=4,
    use_threads=True,
)


# ----------------------------------------------------------------------------
# S3 客户端
# ----------------------------------------------------------------------------
def get_blob_s3_client():
    kwargs = dict(
        s3={"addressing_style": "path"},
        region_name="HB1",
        signature_version=botocore.UNSIGNED,
        connect_timeout=15,
        read_timeout=120,
        retries={"max_attempts": 3},
        # boto3 >=1.36 默认给 PutObject 加 CRC32 校验，body 会以 aws-chunked 分块编码发送。
        # BS3 网关不解这层编码，会把 "100000\r\n" 这样的分块头当成文件内容原样存下来，
        # 导致上传的 mp4 每 1MB 被插一段垃圾、末尾多个 checksum 尾巴，播放器直接打不开。
        # 关掉它，body 才是原始字节。
        request_checksum_calculation="when_required",
        response_checksum_validation="when_required",
    )
    try:
        cfg = Config(**kwargs)
    except TypeError:
        # botocore <1.36 没有这两个参数，本来也不会加 aws-chunked
        kwargs.pop("request_checksum_calculation", None)
        kwargs.pop("response_checksum_validation", None)
        cfg = Config(**kwargs)

    client = boto3.client("s3", endpoint_url=ENDPOINT, config=cfg,
                          use_ssl=ENDPOINT.startswith("https"))

    def _add_service_header(request, **kwargs):
        request.headers.add_header("service", SERVICE_HEADER)

    client.meta.events.register("before-sign.*.*", _add_service_header)
    return client


# ----------------------------------------------------------------------------
# 进度条
# ----------------------------------------------------------------------------
class ProgressPercentage:
    def __init__(self, filename, file_size):
        self._filename = os.path.basename(filename)
        self._size = max(file_size, 1)
        self._seen = 0
        self._lock = threading.Lock()
        self._start = time.time()
        self._last = 0.0

    def __call__(self, bytes_amount):
        with self._lock:
            self._seen += bytes_amount
            now = time.time()
            if now - self._last < 0.5 and self._seen < self._size:
                return
            self._last = now
            elapsed = now - self._start
            pct = self._seen / self._size * 100
            speed = self._seen / (1024 * 1024) / elapsed if elapsed > 0 else 0
            filled = int(30 * self._seen // self._size)
            bar = "#" * filled + "-" * (30 - filled)
            sys.stdout.write(
                "\r%-28s |%s| %5.1f%% | %.1f/%.1fMB | %.1f MB/s"
                % (self._filename[:28], bar, pct,
                   self._seen / (1024 * 1024), self._size / (1024 * 1024), speed))
            sys.stdout.flush()


def clear_line():
    sys.stdout.write("\r" + " " * 110 + "\r")
    sys.stdout.flush()


def human_size(n):
    if n < 1024 * 1024:
        return "%.1f KB" % (n / 1024)
    if n < 1024 * 1024 * 1024:
        return "%.1f MB" % (n / (1024 * 1024))
    return "%.2f GB" % (n / (1024 * 1024 * 1024))


# ----------------------------------------------------------------------------
# 状态记录（避免重复上传）
# ----------------------------------------------------------------------------
def load_state():
    if not os.path.exists(STATE_FILE):
        return {}
    try:
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception:
        return {}


def save_state(state):
    os.makedirs(os.path.dirname(STATE_FILE), exist_ok=True)
    tmp = STATE_FILE + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(state, f, ensure_ascii=False, indent=1)
    os.replace(tmp, STATE_FILE)


# ----------------------------------------------------------------------------
# 扫描 / 过滤
# ----------------------------------------------------------------------------
def should_skip_name(name):
    return any(fnmatch.fnmatch(name, p) for p in SKIP_PATTERNS)


def iter_candidates(root):
    """遍历 root，产出 (本地绝对路径, 相对 root 的 posix 路径)。"""
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = [d for d in dirnames
                       if d not in SKIP_DIRS and not should_skip_name(d)]
        for name in filenames:
            if should_skip_name(name):
                continue
            if os.path.splitext(name)[1].lower() not in ALLOW_EXT:
                continue
            full = os.path.join(dirpath, name)
            rel = os.path.relpath(full, root).replace(os.sep, "/")
            yield full, rel


def build_key(rel_from_local_root):
    return "%s/%s" % (KEY_PREFIX.strip("/"), rel_from_local_root)


def build_url(key):
    return "%s/%s/%s?x-bs-client-force=true&ts=%d" % (
        PUBLIC_HOST, BUCKET, quote(key, safe="/"), int(time.time() * 1000))


def remote_size(client, key):
    try:
        return client.head_object(Bucket=BUCKET, Key=key)["ContentLength"]
    except botocore.exceptions.ClientError as e:
        if e.response["Error"]["Code"] in ("404", "NoSuchKey", "NotFound"):
            return None
        raise


# ----------------------------------------------------------------------------
# 上传
# ----------------------------------------------------------------------------
def upload_one(client, local_path, key, size, quiet=False):
    ctype = mimetypes.guess_type(local_path)[0] or "application/octet-stream"
    cb = None if quiet else ProgressPercentage(local_path, size)
    try:
        client.upload_file(local_path, BUCKET, key,
                           ExtraArgs={"ContentType": ctype},
                           Callback=cb, Config=TRANSFER_CONFIG)
    except Exception as e:
        clear_line()
        print("  [FAIL] %s  -> %s" % (key, e))
        return False, str(e)

    clear_line()

    # 上传成功不代表存对了：核对线上大小，防止 aws-chunked 之类的编码把文件撑大/截断
    try:
        head = client.head_object(Bucket=BUCKET, Key=key)
    except Exception as e:
        print("  [WARN] %s 已上传，但校验失败: %s" % (key, e))
        return True, ""

    got = head["ContentLength"]
    if got != size:
        enc = head.get("ContentEncoding")
        print("  [BAD]  %s 线上大小 %d != 本地 %d%s，文件已损坏"
              % (key, got, size, "（ContentEncoding=%s）" % enc if enc else ""))
        return False, "size mismatch: remote=%d local=%d" % (got, size)

    print("  [OK]   %s  (%s)" % (key, human_size(size)))
    return True, ""


def scan_and_upload(client, root, scan_root, state, args):
    """扫描一轮，返回 (本轮新上传的 [(key, url, local_path)], 因刚写入而跳过的文件数)。

    root 决定 key 怎么算（相对 upload_content），scan_root 决定实际遍历哪个目录，
    这样 --dir 只缩小遍历范围而不会改变线上路径。
    """
    now = time.time()
    todo, waiting = [], 0

    for full, _ in iter_candidates(scan_root):
        try:
            st = os.stat(full)
        except OSError:
            continue
        if now - st.st_mtime < STABLE_SECONDS:
            waiting += 1        # 可能还在写盘，下一轮再看
            continue

        rel = os.path.relpath(full, root).replace(os.sep, "/")
        key = build_key(rel)
        rec = state.get(key)
        if not args.force and not args.repair and rec \
                and rec.get("size") == st.st_size \
                and abs(rec.get("mtime", 0) - st.st_mtime) < 1:
            continue
        todo.append((full, key, st.st_size, st.st_mtime))

    if not todo:
        return [], waiting

    total = sum(t[2] for t in todo)
    print("发现 %d 个待上传文件，共 %s" % (len(todo), human_size(total)))

    uploaded = []
    for full, key, size, mtime in todo:
        if args.dry_run:
            print("  [DRY]  %s  (%s)" % (key, human_size(size)))
            continue

        if not args.force and (args.verify_remote or args.repair):
            if remote_size(client, key) == size:
                print("  [SKIP] 线上已存在且大小一致: %s" % key)
                state[key] = {"size": size, "mtime": mtime,
                              "url": build_url(key), "uploaded_at": time.time()}
                continue
            if args.repair:
                print("  [FIX]  线上文件损坏或缺失，重传: %s" % key)

        ok, _ = upload_one(client, full, key, size)
        if ok:
            url = build_url(key)
            state[key] = {"size": size, "mtime": mtime,
                          "url": url, "uploaded_at": time.time()}
            uploaded.append((key, url, full))
            save_state(state)

    return uploaded, waiting


# ----------------------------------------------------------------------------
# URL 落盘 / 回填
# ----------------------------------------------------------------------------
def write_url_csv(root, state):
    """在每个一级版本目录（0729、0805_test1 …）下写一份 URL 清单。"""
    groups = {}
    prefix = KEY_PREFIX.strip("/") + "/"
    for key, rec in state.items():
        if not key.startswith(prefix):
            continue
        rel = key[len(prefix):]
        top = rel.split("/")[0]
        groups.setdefault(top, []).append((rel, rec))

    for top, items in groups.items():
        ver_dir = os.path.join(root, top)
        if not os.path.isdir(ver_dir):
            continue
        out = os.path.join(ver_dir, top + URL_CSV_SUFFIX)
        with open(out, "w", encoding="utf-8-sig", newline="") as f:
            w = csv.writer(f)
            w.writerow(["filename", "relative_path", "size_bytes", "url"])
            for rel, rec in sorted(items):
                w.writerow([os.path.basename(rel), rel,
                            rec.get("size", ""), rec.get("url", "")])
        print("URL 清单: %s" % out)

        legacy = os.path.join(ver_dir, LEGACY_URL_CSV)
        if os.path.abspath(legacy) != os.path.abspath(out) and os.path.isfile(legacy):
            try:
                os.remove(legacy)
                print("  已清理旧命名: %s" % legacy)
            except OSError as e:
                print("  [WARN] 旧文件删不掉 %s: %s" % (legacy, e))


def fill_xlsx(root, state):
    """把 URL 回填到版本目录下的 xlsx（按文件名匹配 source/target/reference 三列）。"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[WARN] 未安装 openpyxl，跳过 xlsx 回填（pip install openpyxl）")
        return

    # 各版本目录下文件重名（target_video_01.mp4 每天都有），URL 必须按版本目录分开查
    by_ver = {}
    prefix = KEY_PREFIX.strip("/") + "/"
    for key, rec in state.items():
        if not key.startswith(prefix):
            continue
        parts = key[len(prefix):].split("/")
        if len(parts) < 2:
            continue
        by_ver.setdefault(parts[0], {})[parts[-1]] = rec.get("url", "")

    pairs = [("source_video", "source_video_url"),
             ("target_video", "target_video_url"),
             ("reference_image", "reference_image_url")]

    for top in sorted(os.listdir(root)):
        ver_dir = os.path.join(root, top)
        if not os.path.isdir(ver_dir):
            continue
        by_name = by_ver.get(top, {})
        for name in os.listdir(ver_dir):
            if not name.endswith(".xlsx") or name.startswith("~$"):
                continue
            path = os.path.join(ver_dir, name)
            wb = load_workbook(path)
            ws = wb.active
            header = [c.value for c in ws[1]]
            changed = 0
            for src_col, url_col in pairs:
                if src_col not in header or url_col not in header:
                    continue
                si, ui = header.index(src_col) + 1, header.index(url_col) + 1
                for r in range(2, ws.max_row + 1):
                    v = ws.cell(row=r, column=si).value
                    if not v:
                        continue
                    url = by_name.get(os.path.basename(str(v)))
                    if url and ws.cell(row=r, column=ui).value != url:
                        ws.cell(row=r, column=ui).value = url
                        changed += 1
            if changed:
                wb.save(path)
                print("回填 %d 个 URL -> %s" % (changed, path))


# ----------------------------------------------------------------------------
# ----------------------------------------------------------------------------
# 发布：把清单 csv、索引 json、看板页面本身也推上去
#
# 看板一旦托管在 BS3 上，就没有 dashboard_server.py 的 /api/* 接口可用了
# （BS3 只是对象存储，请求 /api/files 会被当成访问名为 api 的 bucket）。
# 所以额外传一份 dashboard_index.json，让页面能自己读出线上有哪些文件。
# ----------------------------------------------------------------------------
def put_file(client, local_path, key, ctype=None):
    ctype = ctype or mimetypes.guess_type(local_path)[0] or "application/octet-stream"
    with open(local_path, "rb") as f:
        body = f.read()
    client.put_object(Bucket=BUCKET, Key=key, Body=body, ContentType=ctype)
    return len(body)


def list_remote(client):
    """线上 相对路径 -> 大小（已去掉 Blender_data/ 前缀）。"""
    prefix = KEY_PREFIX.strip("/") + "/"
    out = {}
    for page in client.get_paginator("list_objects_v2").paginate(Bucket=BUCKET, Prefix=prefix):
        for obj in page.get("Contents", []):
            key = obj["Key"]
            if not key.endswith("/"):
                out[key[len(prefix):]] = obj["Size"]
    return out


def find_dashboard_html():
    exact = os.path.join(BASE_DIR, "视频数据看板.html")
    if os.path.isfile(exact):
        return exact
    pages = [os.path.join(BASE_DIR, n) for n in os.listdir(BASE_DIR)
             if n.lower().endswith(".html")]
    return max(pages, key=os.path.getmtime) if pages else None


def scan_manifests(root):
    """复用看板服务的清单探查（行数 / 类型 / 列名）。"""
    try:
        if BASE_DIR not in sys.path:
            sys.path.insert(0, BASE_DIR)
        import dashboard_server
        return dashboard_server.scan_data_root(root)
    except Exception as e:
        print("  [WARN] 清单探查失败，索引里不带清单: %s: %s" % (type(e).__name__, e))
        return []


def publish_dashboard(client, root, with_page=True):
    prefix = KEY_PREFIX.strip("/") + "/"

    for name in sorted(os.listdir(root)):
        csv_path = os.path.join(root, name, name + URL_CSV_SUFFIX)
        if os.path.isfile(csv_path):
            put_file(client, csv_path, prefix + "%s/%s%s" % (name, name, URL_CSV_SUFFIX),
                     "text/csv; charset=utf-8")

    page_url = ""
    page = find_dashboard_html() if with_page else None
    if page:
        put_file(client, page, prefix + os.path.basename(page), "text/html; charset=utf-8")
        page_url = build_url(prefix + os.path.basename(page))

    remote = list_remote(client)
    files = [{"path": p, "size": s, "url": build_url(prefix + p)}
             for p, s in sorted(remote.items()) if p.lower().endswith(MEDIA_EXT)]
    # 只收录线上确实存在的清单，否则页面点开就是 404；索引自己不算清单
    manifests = [m for m in scan_manifests(root)
                 if m.get("src") in remote and m.get("src") != INDEX_NAME]

    index = {
        "generated_at": time.strftime("%Y-%m-%d %H:%M:%S"),
        "base_url": "%s/%s/%s" % (PUBLIC_HOST, BUCKET, prefix),
        "files": files,
        "manifests": manifests,
    }
    local_index = os.path.join(root, INDEX_NAME)
    with open(local_index, "w", encoding="utf-8") as f:
        json.dump(index, f, ensure_ascii=False, indent=1)
    put_file(client, local_index, prefix + INDEX_NAME, "application/json; charset=utf-8")

    print("线上索引: %d 个媒体文件 / %d 份清单" % (len(files), len(manifests)))
    if page_url:
        print("在线看板: %s" % page_url)
    return len(files), len(manifests)


def apply_overrides(args):
    """命令行 / 指定配置文件 覆盖模块级配置。

    STATE_FILE 是从 LOCAL_ROOT 推出来的，换了根目录必须跟着换，
    否则上传记录会写到旧目录里，下次全当成新文件重传一遍。
    """
    global CONFIG, CONFIG_PATH, LOCAL_ROOT, BUCKET, KEY_PREFIX
    global ENDPOINT, PUBLIC_HOST, SERVICE_HEADER, STATE_FILE

    if args.config:
        if not os.path.isfile(args.config):
            print("[ERR] 配置文件不存在: %s" % args.config)
            sys.exit(1)
        # --root 的默认值是加载 --config 之前的 LOCAL_ROOT，得先记下来才认得出
        # 用户到底有没有显式指定 --root
        untouched = args.root == LOCAL_ROOT
        CONFIG, CONFIG_PATH = load_config(args.config)
        LOCAL_ROOT = CONFIG["local_root"]
        BUCKET = CONFIG["bucket"]
        KEY_PREFIX = CONFIG["key_prefix"]
        ENDPOINT = CONFIG["endpoint"]
        PUBLIC_HOST = CONFIG["public_host"]
        SERVICE_HEADER = CONFIG["service_header"]
        if untouched:
            args.root = LOCAL_ROOT

    if args.bucket:
        BUCKET = args.bucket
    if args.prefix:
        KEY_PREFIX = args.prefix.strip("/")

    args.root = os.path.abspath(os.path.expanduser(args.root))
    LOCAL_ROOT = args.root
    STATE_FILE = os.path.join(LOCAL_ROOT, ".blob_upload_state.json")


def main():
    p = argparse.ArgumentParser(description="upload_content 自动上传到 BS3")
    p.add_argument("--root", default=LOCAL_ROOT, help="本地监听根目录")
    p.add_argument("--dir", default=None, help="只处理根目录下的某个子目录，如 0805")
    p.add_argument("--once", action="store_true", help="只扫描上传一次")
    p.add_argument("--watch", action="store_true", help="常驻监听")
    p.add_argument("--interval", type=int, default=WATCH_INTERVAL, help="监听扫描间隔（秒）")
    p.add_argument("--dry-run", action="store_true", help="只打印不上传")
    p.add_argument("--force", action="store_true", help="忽略本地记录，全部重传")
    p.add_argument("--verify-remote", action="store_true",
                   help="上传前 head_object 核对线上是否已有同样大小的对象")
    p.add_argument("--repair", action="store_true",
                   help="逐个核对线上大小，只重传损坏或缺失的文件（修 aws-chunked 损坏用这个）")
    p.add_argument("--fill-xlsx", action="store_true", help="把 URL 回填进版本目录的 xlsx")
    p.add_argument("--write-urls", action="store_true",
                   help="不上传，只根据已有记录重新生成各批次的 URL 清单 csv")
    p.add_argument("--publish", action="store_true",
                   help="不扫描素材，只把清单 csv、索引 json、看板页面推到线上")
    p.add_argument("--config", default=None,
                   help="指定 blob_config.json，默认找脚本旁边和当前目录")
    p.add_argument("--prefix", default=None,
                   help="线上前缀，覆盖配置里的 key_prefix")
    p.add_argument("--bucket", default=None, help="桶名，覆盖配置里的 bucket")
    p.add_argument("--show-config", action="store_true",
                   help="打印当前生效的配置就退出，排查用")
    args = p.parse_args()

    apply_overrides(args)

    if args.show_config:
        print("配置文件 : %s" % (CONFIG_PATH or "（没找到，用的默认值）"))
        print("本地根目录: %s   存在=%s" % (LOCAL_ROOT, os.path.isdir(LOCAL_ROOT)))
        print("线上位置 : %s/%s/%s/" % (PUBLIC_HOST, BUCKET, KEY_PREFIX))
        print("上传记录 : %s" % STATE_FILE)
        return 0

    if not args.once and not args.watch:
        args.once = True

    root = args.root
    scan_root = os.path.join(root, args.dir) if args.dir else root
    if not os.path.isdir(scan_root):
        print("[ERR] 目录不存在: %s" % scan_root)
        return 1

    try:
        sys.stdout.reconfigure(encoding="utf-8")
    except Exception:
        pass

    state = load_state()

    if args.write_urls:
        if not state:
            print("[ERR] 还没有上传记录（%s 不存在）" % STATE_FILE)
            return 1
        write_url_csv(root, state)
        if args.fill_xlsx:
            fill_xlsx(root, state)
        return 0

    client = get_blob_s3_client()

    if args.publish:
        if state:
            write_url_csv(root, state)
        publish_dashboard(client, root)
        return 0

    print("=" * 78)
    print("监听目录 : %s" % scan_root)
    print("线上位置 : %s/%s/%s/" % (PUBLIC_HOST, BUCKET, KEY_PREFIX))
    print("模式     : %s%s" % ("watch" if args.watch else "once",
                               "（dry-run）" if args.dry_run else ""))
    print("=" * 78)

    def run():
        uploaded, waiting = scan_and_upload(client, root, scan_root, state, args)
        if uploaded and not args.dry_run:
            print("-" * 78)
            for key, url, _ in uploaded:
                print(url)
            write_url_csv(root, state)
            if args.fill_xlsx:
                fill_xlsx(root, state)
            try:
                publish_dashboard(client, root)
            except Exception as e:
                print("[WARN] 发布线上索引失败: %s: %s" % (type(e).__name__, e))
        return uploaded, waiting

    if args.once:
        uploaded, waiting = run()
        if not uploaded and not args.dry_run:
            print("没有需要上传的新文件%s" %
                  ("（%d 个文件刚写入，稍后再试）" % waiting if waiting else ""))
        return 0

    print("已进入监听模式，Ctrl+C 退出。把文件丢进上面的目录即可自动上传。")
    # 常驻进程最容易踩的坑：脚本改了但进程还跑着旧代码，行为和预期对不上。
    # 发现自身源码变了就原地重启，保证跑的永远是最新逻辑。
    src_mtime = os.path.getmtime(__file__)
    try:
        while True:
            run()
            time.sleep(args.interval)
            try:
                now_mtime = os.path.getmtime(__file__)
            except OSError:
                continue
            if now_mtime != src_mtime:
                print("\n检测到脚本已更新，正在重启监听以加载新代码…")
                sys.stdout.flush()
                os.execv(sys.executable, [sys.executable] + sys.argv)
    except KeyboardInterrupt:
        print("\n已停止监听。")
    return 0


if __name__ == "__main__":
    sys.exit(main())
